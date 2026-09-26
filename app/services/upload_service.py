import os
import re
import time
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.repositories.inventory_repository import InventoryRepository
from app.repositories.upload_history_repository import UploadHistoryRepository
from app.services.set_calculation import calculate_set_inventory

REQUIRED_HEADER_GROUPS = (
    ("item", "item (child code)", "main code", "child code"),
    ("description", "descriptions"),
    ("location",),
    ("lpn",),
    ("on hand", "onhand", "qty", "quantity", "available", "available qty"),
)

REQUIRED_COLUMN_MESSAGE = "Item, Description, Location, LPN, Quantity"


class InventoryUploadService:
    """Service for validating and replacing inventory records from uploaded Excel files."""

    def __init__(self):
        self.inventory_repository = InventoryRepository()
        self.upload_history_repository = UploadHistoryRepository()

    @staticmethod
    def _safe_value(value: Any) -> str:
        if value is None or value == "":
            return ""
        return str(value).strip()

    @staticmethod
    def _safe_numeric(value: Any, default: float = 0.0) -> float:
        if value is None or value == "":
            return default
        try:
            return float(value)
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _read_quantity(record: dict[str, Any]) -> float:
        """Use the uploaded Qty field as the source of truth and only fall back to other quantity columns if Qty is absent."""
        for key in ("Qty", "Quantity", "On Hand", "Onhand", "Available Qty", "Available", "available_qty"):
            if record.get(key) is not None:
                try:
                    return float(record.get(key))
                except (TypeError, ValueError):
                    return 0.0
        return 0.0

    @staticmethod
    def _normalize_header(value: Any) -> str:
        return re.sub(r"\s+", " ", str(value).strip()).lower()

    @classmethod
    def _canonicalize_rows(cls, raw_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        if not raw_rows:
            return []

        canonicalized_rows: list[dict[str, Any]] = []
        alias_map = {
            "item": "Item (Child code)",
            "item (child code)": "Item (Child code)",
            "child code": "Child Code",
            "main code": "Main Code",
            "description": "description",
            "descriptions": "description",
            "lpn": "LPN",
            "location": "Location",
            "onhand": "On Hand",
            "on hand": "On Hand",
            "qty": "Qty",
            "quantity": "Qty",
            "available": "Available",
            "available qty": "Available Qty",
            "set qty": "Set Qty",
            "loose qty": "Loose Qty",
        }

        for row in raw_rows:
            mapped_row: dict[str, Any] = {}
            for key, value in row.items():
                normalized_key = cls._normalize_header(key)
                destination_key = alias_map.get(normalized_key, str(key).strip())
                mapped_row[destination_key] = value
            canonicalized_rows.append(mapped_row)

        return canonicalized_rows

    @staticmethod
    def _read_excel_rows(file_path: str) -> list[dict[str, Any]]:
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                return []

            headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
            data_rows: list[dict[str, Any]] = []
            for row in rows[1:]:
                record: dict[str, Any] = {}
                for idx, header in enumerate(headers):
                    if idx < len(row):
                        record[header] = row[idx]
                data_rows.append(record)
            return data_rows
        finally:
            workbook.close()

    @staticmethod
    def _has_required_columns(normalized_columns: set[str]) -> bool:
        return all(
            any(alias in normalized_columns for alias in required_group)
            for required_group in REQUIRED_HEADER_GROUPS
        )

    def _prepare_valid_rows(self, raw_rows: list[dict[str, Any]]) -> tuple[list[dict], int]:
        canonical_rows = self._canonicalize_rows(raw_rows)
        columns = set().union(*(record.keys() for record in canonical_rows)) if canonical_rows else set()
        normalized_columns = {self._normalize_header(column) for column in columns}

        if canonical_rows and self._has_required_columns(normalized_columns):
            computed_rows = calculate_set_inventory(canonical_rows)
            return [
                {
                    "main_code": row["Main Code"],
                    "child_code": row["Child Code"],
                    "description": row["Description"],
                    "available_qty": row.get("Available Qty", self._read_quantity(row)),
                    "set_qty": row.get("Set Qty", 0),
                    "loose_qty": row.get("Loose Qty", 0),
                }
                for row in computed_rows
            ], 0

        valid_rows: list[dict] = []
        failed_count = 0

        for record in canonical_rows:
            main_code = self._safe_value(record.get("Main Code"))
            child_code = self._safe_value(record.get("Child Code"))
            description = self._safe_value(record.get("Description"))

            if not all([main_code, child_code, description]):
                failed_count += 1
                continue

            valid_rows.append(
                {
                    "main_code": main_code,
                    "child_code": child_code,
                    "description": description,
                    "available_qty": self._read_quantity(record),
                    "set_qty": self._safe_numeric(record.get("Set Qty")),
                    "loose_qty": self._safe_numeric(record.get("Loose Qty")),
                }
            )

        return valid_rows, failed_count

    def process_upload(self, file_path: str, uploaded_by: str) -> dict:
        """Load and replace master inventory based on a compliant Excel file."""
        start_time = time.perf_counter()
        db: Session = SessionLocal()
        batch = None

        try:
            if not os.path.exists(file_path):
                raise FileNotFoundError(f"Uploaded file not found: {file_path}")

            raw_rows = self._read_excel_rows(file_path)
            if not raw_rows:
                raise ValueError("Excel file is empty")

            canonical_rows = self._canonicalize_rows(raw_rows)
            columns = set().union(*(record.keys() for record in canonical_rows)) if canonical_rows else set()
            normalized_columns = {self._normalize_header(column) for column in columns}

            if not self._has_required_columns(normalized_columns):
                raise ValueError(f"Missing required columns: {REQUIRED_COLUMN_MESSAGE}")

            valid_rows, failed_count = self._prepare_valid_rows(canonical_rows)
            total_records = len(raw_rows)
            successful_records = len(valid_rows)

            batch = self.upload_history_repository.create_batch(
                db,
                file_name=os.path.basename(file_path),
                uploaded_by=uploaded_by,
                total_records=0,
                success_records=0,
                failed_records=0,
                status="in_progress",
            )
            db.commit()

            with db.begin():
                self.inventory_repository.replace_inventory_records(db, valid_rows, batch.batch_id)

            batch.status = "completed"
            batch.total_records = total_records
            batch.success_records = successful_records
            batch.failed_records = failed_count
            db.add(batch)
            db.commit()

            elapsed_seconds = time.perf_counter() - start_time
            return {
                "total_records": total_records,
                "successful_records": successful_records,
                "failed_records": failed_count,
                "upload_time": f"{elapsed_seconds:.2f}s",
                "batch_id": batch.batch_id,
                "status": "completed",
            }
        except Exception:
            db.rollback()
            if batch is not None:
                batch.status = "failed"
                batch.total_records = 0
                batch.success_records = 0
                batch.failed_records = 0
                db.add(batch)
                db.commit()
            raise
        finally:
            db.close()

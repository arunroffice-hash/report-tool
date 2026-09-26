from io import BytesIO
from typing import Dict, Tuple
import logging

from openpyxl import load_workbook
from openpyxl.writer.excel import save_workbook
from openpyxl import Workbook
import itertools


def _normalize_header_str(cell: object) -> str:
    if cell is None:
        return ""
    s = str(cell)
    # normalize common invisible characters and whitespace
    s = s.replace("\u00A0", " ")  # no-break space
    s = s.replace("\ufeff", "")  # BOM
    return s.strip().lower()


def _norm(val: object) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _key(item: str, location: str) -> Tuple[str, str]:
    return (item.strip().lower(), location.strip().lower())


def _read_sheet_map(fileobj) -> Dict[Tuple[str, str], dict]:
    """Read an Excel file (file-like) in streaming/read-only mode and
    return a mapping of (item, location) -> {item, description, location, total}.
    Duplicate keys will be summed (Total aggregated).
    """
    wb = load_workbook(filename=fileobj, read_only=True, data_only=True)
    ws = wb.active

    logger = logging.getLogger("app.services.variance_service")

    rows_iter = ws.iter_rows(values_only=True)

    # Buffer the first N rows and try to locate the header row robustly
    buffered = []
    header_row_idx = None
    headers = None
    MAX_HEADER_SCAN = 20
    for i in range(MAX_HEADER_SCAN):
        try:
            r = next(rows_iter)
        except StopIteration:
            break
        buffered.append(r)
        normalized = [_normalize_header_str(c) for c in r]
        # Heuristic: header row must contain 'item' and at least one of 'location' or 'total'
        if "item" in normalized and ("location" in normalized or "total" in normalized):
            header_row_idx = i
            headers = normalized
            break

    if headers is None:
        # fallback to first buffered row or empty
        if buffered:
            headers = [_normalize_header_str(c) for c in buffered[0]]
            header_row_idx = 0
        else:
            return {}

    logger.info("Detected sheet headers: %s (row=%s)", headers, header_row_idx)
    print("[variance_service] Detected sheet headers:", headers, flush=True)
    # also persist to a debug file for investigation when console logs are not visible
    try:
        with open("tmp_variance_debug.log", "a", encoding="utf-8") as df:
            df.write(f"HEADERS: {headers} ROW={header_row_idx}\n")
    except Exception:
        pass

    # Build an iterator that yields rows starting after the detected header
    remaining_iter = itertools.chain(buffered[header_row_idx + 1 :], rows_iter)

    # identify column indices
    idx_item = None
    idx_description = None
    idx_location = None
    idx_total = None
    idx_division = None
    for i, h in enumerate(headers):
        if h == "item":
            idx_item = i
        elif h == "description":
            idx_description = i
        elif h == "location":
            idx_location = i
        elif h == "total":
            idx_total = i
        elif h == "division":
            idx_division = i

    result = {}
    sample_rows = []
    for row_idx, row in enumerate(remaining_iter, start=1):
        if row_idx <= 3:
            sample_rows.append(row)
        # handle rows shorter than header length
        row_len = len(row) if row is not None else 0
        item = _norm(row[idx_item]) if idx_item is not None and idx_item < row_len else ""
        description = _norm(row[idx_description]) if idx_description is not None and idx_description < row_len else ""
        location = _norm(row[idx_location]) if idx_location is not None and idx_location < row_len else ""
        total_raw = row[idx_total] if idx_total is not None and idx_total < row_len else 0
        try:
            total = float(total_raw) if total_raw is not None and total_raw != "" else 0.0
        except Exception:
            total = 0.0

        division = _norm(row[idx_division]) if idx_division is not None and idx_division < row_len else ""

        if item == "" and location == "":
            continue

        k = _key(item, location)
        if k in result:
            result[k]["total"] += total
        else:
            result[k] = {
                "item": item,
                "description": description,
                "location": location,
                "total": total,
                "division": division,
            }

    wb.close()
    if sample_rows:
        logger.info("Sample rows (first %d): %s", len(sample_rows), sample_rows)
        print("[variance_service] Sample rows:", sample_rows, flush=True)
        try:
            with open("tmp_variance_debug.log", "a", encoding="utf-8") as df:
                df.write(f"SAMPLES: {sample_rows}\n")
        except Exception:
            pass
    return result


def generate_variance_workbook(old_fileobj, new_fileobj) -> BytesIO:
    """Given two file-like objects for old and new reports, produce an
    Excel workbook as a BytesIO buffer containing the comparison and two
    extra sheets listing items present only in one sheet.
    """
    old_map = _read_sheet_map(old_fileobj)
    new_map = _read_sheet_map(new_fileobj)

    all_keys = set(old_map.keys()) | set(new_map.keys())

    wb = Workbook(write_only=True)

    main_ws = wb.create_sheet(title="Variance Report")
    main_ws.append(["Item", "Description", "Location", "Division", "Old Total", "New Total", "Variance", "Remarks"])

    old_only_ws = wb.create_sheet(title="In Old not in New")
    old_only_ws.append(["Item", "Description", "Location", "Division", "Total"])

    new_only_ws = wb.create_sheet(title="In New not in Old")
    new_only_ws.append(["Item", "Description", "Location", "Division", "Total"])

    for k in sorted(all_keys):
        old_row = old_map.get(k)
        new_row = new_map.get(k)

        item = new_row["item"] if new_row else (old_row["item"] if old_row else "")
        description = new_row["description"] if new_row else (old_row["description"] if old_row else "")
        location = new_row["location"] if new_row else (old_row["location"] if old_row else "")
        division = new_row.get("division", "") if new_row else (old_row.get("division", "") if old_row else "")

        old_total = old_row["total"] if old_row else 0.0
        new_total = new_row["total"] if new_row else 0.0
        variance = new_total - old_total
        remarks = "Match" if abs(variance) < 1e-9 else "Mismatch"
        main_ws.append([item, description, location, division, old_total, new_total, variance, remarks])

        if old_row and not new_row:
            old_only_ws.append([old_row["item"], old_row["description"], old_row["location"], old_row.get("division", ""), old_row["total"]])
        if new_row and not old_row:
            new_only_ws.append([new_row["item"], new_row["description"], new_row["location"], new_row.get("division", ""), new_row["total"]])

    # save to BytesIO
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio

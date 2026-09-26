from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.services.variance_service import generate_variance_workbook


def _make_excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Description", "Location", "Total"])
    for r in rows:
        ws.append(r)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def test_generate_variance_workbook_basic():
    old_rows = [
        ["A1", "Item A1", "Loc1", 10],
        ["B2", "Item B2", "Loc2", 5],
    ]
    new_rows = [
        ["A1", "Item A1", "Loc1", 12],
        ["C3", "Item C3", "Loc3", 7],
    ]

    old_bio = _make_excel(old_rows)
    # include Division column in new file
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Description", "Location", "Division", "Total"])
    for r in [["A1", "Item A1", "Loc1", "Accessory", 12], ["C3", "Item C3", "Loc3", "Furniture", 7]]:
        ws.append(r)
    new_bio = BytesIO()
    wb.save(new_bio)
    new_bio.seek(0)

    out = generate_variance_workbook(old_bio, new_bio)
    assert out is not None

    book = load_workbook(out)
    assert "Variance Report" in book.sheetnames
    assert "In Old not in New" in book.sheetnames
    assert "In New not in Old" in book.sheetnames

    main = book["Variance Report"]
    main_headers = [c.value for c in main[1]]
    assert "Division" in main_headers

    old_only = book["In Old not in New"]
    old_headers = [c.value for c in old_only[1]]
    assert "Division" in old_headers

    new_only = book["In New not in Old"]
    new_headers = [c.value for c in new_only[1]]
    assert "Division" in new_headers

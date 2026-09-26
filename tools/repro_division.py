from io import BytesIO
import os
import sys

# Ensure project root is on sys.path for imports when run as a script
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from openpyxl import Workbook, load_workbook

from app.services.variance_service import generate_variance_workbook


def make_excel_with_headers(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def main():
    # Old file (no Division)
    old_headers = ["Item", "Description", "Location", "Old Total"]
    old_rows = [
        ["1000", "Desc A", "Loc1", 10],
        ["1001", "Desc B", "Loc2", 5],
    ]
    old_bio = make_excel_with_headers(old_headers, old_rows)

    # New file has Division as first column and 'Main' between
    new_headers = ["Division", "Main", "Item", "Description", "Location", "Total"]
    new_rows = [
        ["Accessory", "M1", "1000", "Desc A", "Loc1", 12],
        ["Furniture", "M2", "1002", "Desc C", "Loc3", 7],
    ]
    new_bio = make_excel_with_headers(new_headers, new_rows)

    out = generate_variance_workbook(old_bio, new_bio)

    book = load_workbook(out)
    print("Generated sheets:", book.sheetnames)
    main = book["Variance Report"]
    header = [c.value for c in main[1]]
    print("Main header:", header)


if __name__ == "__main__":
    main()
